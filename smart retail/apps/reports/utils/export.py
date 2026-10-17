import csv
import io
from django.http import HttpResponse


def export_to_csv(rows, columns, filename="report.csv"):
    """`columns`: list of (field_key, header_label). `rows`: list of dicts."""
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)
    writer.writerow([label for _, label in columns])
    for row in rows:
        writer.writerow([row.get(key, "") for key, _ in columns])
    return response


def export_to_excel(rows, columns, filename="report.xlsx", sheet_title="Report", header_lines=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Optional title block (company name, report title, period) above the table
    for i, text in enumerate(header_lines or []):
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=14 if i == 0 else 11)
    if header_lines:
        ws.append([])

    head_row = ws.max_row + 1
    ws.append([label for _, label in columns])
    for cell in ws[head_row]:
        cell.font = Font(bold=True)

    statement = any("kind" in r for r in rows)   # P&L statement style rows
    numeric_keys = {"amount", "percent"}
    top_rule = Border(top=Side(style="thin"))
    for row in rows:
        ws.append([row.get(key, "") for key, _ in columns])
        if not statement:
            continue
        r = ws.max_row
        for ci, (key, _) in enumerate(columns, start=1):
            c = ws.cell(row=r, column=ci)
            if row.get("_bold"):
                c.font = Font(bold=True, underline="single" if row.get("kind") == "heading_u" and key == "description" else None)
            if key in numeric_keys:
                c.alignment = Alignment(horizontal="right")
                if row.get("kind") == "subtotal":
                    c.border = top_rule

    for ci in range(1, len(columns) + 1):
        letter = ws.cell(row=head_row, column=ci).column_letter
        vals = [ws.cell(row=r, column=ci).value for r in range(head_row, ws.max_row + 1)]
        max_len = max((len(str(v)) for v in vals if v is not None), default=10)
        ws.column_dimensions[letter].width = min(max_len + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_to_pdf(rows, columns, filename="report.pdf", title="Report", header_lines=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    statement = any("kind" in r for r in rows)   # P&L statement: portrait, ledger-style table
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4 if statement else landscape(A4),
                             leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                             topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()

    if header_lines:
        elements = [Paragraph(header_lines[0], styles["Title"])]
        for text in header_lines[1:]:
            elements.append(Paragraph(text, styles["Heading4"] if text.isupper() else styles["Normal"]))
        elements.append(Spacer(1, 12))
    else:
        elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    header = [label for _, label in columns]
    data = [header] + [[str(row.get(key, "")) for key, _ in columns] for row in rows]

    if statement:
        # ledger-style widths: code | description (wide) | amount | %
        widths = {"code": 2.6 * cm, "description": 8.6 * cm, "amount": 3.6 * cm, "percent": 2.2 * cm}
        table = Table(data, colWidths=[widths.get(k, 3 * cm) for k, _ in columns], repeatRows=1, hAlign="LEFT")
    else:
        table = Table(data, repeatRows=1)
    if statement:
        keys = [k for k, _ in columns]
        cmds = [
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]
        for ci, k in enumerate(keys):
            if k in ("amount", "percent"):
                cmds.append(("ALIGN", (ci, 0), (ci, -1), "RIGHT"))
        di = keys.index("description") if "description" in keys else 0
        for ri, row in enumerate(rows, start=1):
            if row.get("_bold"):
                cmds.append(("FONTNAME", (0, ri), (-1, ri), "Helvetica-Bold"))
            if row.get("kind") == "subtotal":
                for ci, k in enumerate(keys):
                    if k in ("amount", "percent"):
                        cmds.append(("LINEABOVE", (ci, ri), (ci, ri), 0.6, colors.black))
            if row.get("kind") == "line":
                cmds.append(("LEFTPADDING", (di, ri), (di, ri), 6))
        table.setStyle(TableStyle(cmds))
    else:
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4f46e5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
    elements.append(table)
    doc.build(elements)

    buffer.seek(0)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def export_response(request, rows, columns, base_filename, title="Report", header_lines=None):
    """Dispatches to csv/excel/pdf based on ?export= query param. Defaults to a normal
    JSON response (return None) if unset. NOTE: intentionally NOT named 'format' —
    that query param name is reserved by DRF for content-negotiation and colliding
    with it causes DRF to raise Http404 for any value it doesn't recognize as a renderer."""
    fmt = request.query_params.get("export", "").lower()
    if fmt == "csv":
        return export_to_csv(rows, columns, f"{base_filename}.csv")
    if fmt in ("excel", "xlsx"):
        return export_to_excel(rows, columns, f"{base_filename}.xlsx", sheet_title=title, header_lines=header_lines)
    if fmt == "pdf":
        return export_to_pdf(rows, columns, f"{base_filename}.pdf", title=title, header_lines=header_lines)
    return None  # caller falls back to a normal DRF Response (JSON)